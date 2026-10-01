from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


def extract_resume_text(filename, content):
    extension = Path(filename).suffix.lower()
    try:
        if extension == ".txt":
            text = content.decode("utf-8-sig")
        elif extension == ".pdf":
            reader = PdfReader(BytesIO(content), strict=True)
            if reader.is_encrypted:
                raise ValueError("This PDF is password-protected. Upload an unlocked copy.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages[:20])
        elif extension == ".docx":
            document = Document(BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            table_cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            text = "\n".join(paragraphs + table_cells)
        else:
            raise ValueError("That file type is not supported. Choose PDF, DOCX, or TXT.")
    except UnicodeDecodeError as error:
        raise ValueError("This text file isn't UTF-8 encoded. Save it as UTF-8 and try again.") from error
    except ValueError:
        raise
    except Exception as error:
        raise ValueError("We couldn't read this file. Check that it opens correctly, then try again.") from error
    return text[:30000]