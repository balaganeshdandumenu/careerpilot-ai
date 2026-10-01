import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from docx import Document
from google.genai import types

from app import create_app
from config.config import Config
from services import rate_limit_service
from services import session_store
from services.gemini_service import generate_json
from utils.file_parser import extract_resume_text


class TestConfig(Config):
    TESTING = True
    WTF_CSRF_ENABLED = False
    SECRET_KEY = "test-secret"
    GEMINI_API_KEY = None


class CareerPilotTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.original_database = session_store.DATABASE
        self.original_rate_database = rate_limit_service.DATABASE
        session_store.DATABASE = Path(self.temp_dir.name) / "sessions.sqlite3"
        rate_limit_service.DATABASE = Path(self.temp_dir.name) / "limits" / "limits.sqlite3"
        self.app = create_app(TestConfig)
        self.client = self.app.test_client()

    def tearDown(self):
        session_store.DATABASE = self.original_database
        rate_limit_service.DATABASE = self.original_rate_database
        self.temp_dir.cleanup()

    def test_career_plan_has_thirty_days_and_fits_four_weeks(self):
        response = self.client.post("/career/generate", data={
            "name": "Alex",
            "target_role": "ai-ml-engineer",
            "experience_level": "student",
            "skills": ["Python", "SQL"],
            "hours_per_day": "1.5",
            "days_per_week": "5",
            "preferences": ["projects"],
        }, follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"YOUR CAREER SNAPSHOT", response.data)
        roadmap = self.client.get("/career/roadmap")
        self.assertIn(b"DAY 30", roadmap.data)
        self.assertIn(b"0/9 complete", roadmap.data)
        completed = self.client.post("/career/roadmap/complete", data={"day": "1", "return_to": "dashboard"}, follow_redirects=True)
        self.assertEqual(completed.status_code, 200)
        self.assertIn(b"1 day learning streak", completed.data)

    def test_resume_upload_compares_role_and_updates_dashboard(self):
        response = self.client.post("/resume/analyze", data={
            "target_role": "ai-ml-engineer",
            "resume": (BytesIO(b"Alex Candidate\nSkills\nPython\nProjects\nBuilt a Flask model API using SQL and Git."), "resume.txt"),
        }, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"ESTIMATED ATS COMPATIBILITY", response.data)
        self.assertIn(b"Not found in this text", response.data)
        reloaded = self.client.get("/resume/report")
        self.assertEqual(reloaded.status_code, 200)
        self.assertIn(b"ESTIMATED ATS COMPATIBILITY", reloaded.data)

    def test_unsupported_resume_type_has_clear_error(self):
        response = self.client.post("/resume/analyze", data={
            "target_role": "ai-ml-engineer",
            "resume": (BytesIO(b"not a resume"), "resume.exe"),
        }, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 400)
        self.assertIn(b"That file type is not supported", response.data)

    def test_docx_table_text_is_extracted(self):
        document = Document()
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Python · Pandas"
        content = BytesIO()
        document.save(content)
        text = extract_resume_text("resume.docx", content.getvalue())
        self.assertIn("Python", text)
        self.assertIn("Pandas", text)

    def test_rate_limit_blocks_requests_after_threshold(self):
        self.assertTrue(rate_limit_service._allow_request("unit-test", "hashed-client", 2, 60))
        self.assertTrue(rate_limit_service._allow_request("unit-test", "hashed-client", 2, 60))
        self.assertFalse(rate_limit_service._allow_request("unit-test", "hashed-client", 2, 60))

    def test_gemini_malformed_response_gets_one_repair_attempt(self):
        client = MagicMock()
        client.models.generate_content.side_effect = [
            SimpleNamespace(text="not json"),
            SimpleNamespace(text='{"summary":"A grounded summary"}'),
        ]
        schema = types.Schema(
            type=types.Type.OBJECT,
            properties={"summary": types.Schema(type=types.Type.STRING)},
        )
        with patch("services.gemini_service.genai.Client", return_value=client):
            result = generate_json("test-key", "profile facts", schema)
        self.assertEqual(result["summary"], "A grounded summary")
        self.assertEqual(client.models.generate_content.call_count, 2)


if __name__ == "__main__":
    unittest.main()